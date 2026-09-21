from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import BinaryIO

from django.db import transaction
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from rest_framework.exceptions import ValidationError

from apps.orders.models import Order, OrderStatus
from apps.orders.services.change_order_status import change_order_status

HEADER_ORDER_NUMBER = "شماره سفارش"
HEADER_FIRST_NAME = "نام"
HEADER_LAST_NAME = "نام خانوادگی"
HEADER_TRACKING = "کد رهگیری"
HEADER_PHONE = "شماره موبایل"

HEADERS = (
    HEADER_ORDER_NUMBER,
    HEADER_FIRST_NAME,
    HEADER_LAST_NAME,
    HEADER_TRACKING,
    HEADER_PHONE,
)

REQUIRED_HEADERS = (
    HEADER_ORDER_NUMBER,
    HEADER_FIRST_NAME,
    HEADER_LAST_NAME,
    HEADER_TRACKING,
)

_DIGIT_TRANSLATE = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)
_LETTER_RE = re.compile(r"[A-Za-z\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]")
_TRACKING_MIN_LEN = 6
_TRACKING_MAX_LEN = 40


@dataclass
class RowError:
    row_number: int
    order_number: str
    message: str


@dataclass
class ImportResult:
    success_count: int = 0
    errors: list[RowError] = field(default_factory=list)


def normalize_persian_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = " ".join(text.split())
    text = text.replace("ي", "ی").replace("ك", "ک")
    return text


def normalize_tracking_digits(value: str) -> str:
    text = normalize_persian_text(value).translate(_DIGIT_TRANSLATE)
    text = text.replace("ـ", "").replace("-", "").replace("_", "")
    text = re.sub(r"\s+", "", text)
    return text


def validate_tracking_code(raw_value) -> str:
    """Accept digit-only tracking codes (Persian digits OK). Reject letters."""
    text = normalize_persian_text(raw_value)
    if not text:
        raise ValidationError("کد رهگیری خالی است.")

    # Convert Eastern digits first so they are not treated as letters.
    ascii_ready = text.translate(_DIGIT_TRANSLATE)
    if _LETTER_RE.search(ascii_ready):
        raise ValidationError(
            "کد رهگیری نامعتبر است: شامل حرف است. فقط عدد وارد کنید."
        )

    digits = normalize_tracking_digits(text)
    if not digits:
        raise ValidationError("کد رهگیری خالی است.")

    if not digits.isdigit():
        raise ValidationError(
            "کد رهگیری نامعتبر است: فقط عدد مجاز است "
            "(بدون حرف، فاصله یا کاراکتر خاص)."
        )

    if len(digits) < _TRACKING_MIN_LEN:
        raise ValidationError(
            f"کد رهگیری کوتاه است؛ حداقل {_TRACKING_MIN_LEN} رقم لازم است."
        )

    if len(digits) > _TRACKING_MAX_LEN:
        raise ValidationError(
            f"کد رهگیری بیش از حد طولانی است؛ حداکثر {_TRACKING_MAX_LEN} رقم."
        )

    return digits


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_map(header_row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = normalize_persian_text(cell)
        if key:
            mapping[key] = idx
    return mapping


def _order_lookup_key(order: Order) -> str:
    if order.public_number:
        return order.public_number
    return str(order.pk)


def find_order_by_number(order_number: str) -> Order | None:
    raw = normalize_persian_text(order_number)
    if not raw:
        return None

    order = (
        Order.objects.select_related("user", "shipping_method")
        .filter(public_number=raw)
        .first()
    )
    if order:
        return order

    if raw.isdigit():
        return (
            Order.objects.select_related("user", "shipping_method")
            .filter(pk=int(raw))
            .first()
        )
    return None


def names_match(*, order: Order, first_name: str, last_name: str) -> bool:
    user = order.user
    return (
        normalize_persian_text(user.first_name) == normalize_persian_text(first_name)
        and normalize_persian_text(user.last_name) == normalize_persian_text(last_name)
    )


def _apply_tracking_and_ship(
    *,
    order: Order,
    tracking_code: str,
    changed_by,
) -> None:
    update_fields = ["tracking_code", "updated_at"]
    order.tracking_code = tracking_code

    if not (order.carrier or "").strip():
        method = getattr(order, "shipping_method", None)
        if method is None and order.shipping_method_id:
            from apps.shipping.models import ShippingMethod

            method = ShippingMethod.objects.filter(
                pk=order.shipping_method_id,
            ).first()
        if method and method.carrier:
            order.carrier = method.carrier
            update_fields.append("carrier")

    order.save(update_fields=update_fields)

    if order.status == OrderStatus.SHIPPED:
        return

    if order.status != OrderStatus.PREPARING:
        raise ValidationError(
            "وضعیت سفارش باید «در حال آماده‌سازی» باشد "
            f"(وضعیت فعلی: {order.get_status_display()})."
        )

    change_order_status(
        order=order,
        new_status=OrderStatus.SHIPPED,
        changed_by=changed_by,
        reason="Bulk tracking Excel import.",
    )


def import_tracking_from_workbook(
    *,
    file_obj: BinaryIO,
    changed_by,
) -> ImportResult:
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(
            _("فایل اکسل نامعتبر است: %(error)s") % {"error": str(exc)}
        ) from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValidationError(_("فایل اکسل خالی است.")) from exc

    headers = _header_map(header_row)
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValidationError(
            _("ستون‌های الزامی موجود نیست: %(cols)s")
            % {"cols": "، ".join(missing)}
        )

    result = ImportResult()

    for excel_row_index, row in enumerate(rows, start=2):
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def col(name: str) -> str:
            idx = headers[name]
            if idx >= len(row):
                return ""
            return _cell_str(row[idx])

        order_number = col(HEADER_ORDER_NUMBER)
        first_name = col(HEADER_FIRST_NAME)
        last_name = col(HEADER_LAST_NAME)
        tracking_raw = col(HEADER_TRACKING)

        if not order_number and not tracking_raw and not first_name and not last_name:
            continue

        if not order_number:
            result.errors.append(
                RowError(
                    row_number=excel_row_index,
                    order_number="",
                    message="شماره سفارش خالی است.",
                )
            )
            continue

        try:
            tracking_code = validate_tracking_code(tracking_raw)
        except ValidationError as exc:
            detail = exc.detail
            if isinstance(detail, list):
                message = "; ".join(str(item) for item in detail)
            else:
                message = str(detail)
            result.errors.append(
                RowError(
                    row_number=excel_row_index,
                    order_number=order_number,
                    message=message,
                )
            )
            continue

        if not first_name or not last_name:
            result.errors.append(
                RowError(
                    row_number=excel_row_index,
                    order_number=order_number,
                    message="نام و نام خانوادگی الزامی است.",
                )
            )
            continue

        try:
            with transaction.atomic():
                order = find_order_by_number(order_number)
                if order is None:
                    raise ValidationError("سفارش پیدا نشد.")

                if not names_match(
                    order=order,
                    first_name=first_name,
                    last_name=last_name,
                ):
                    raise ValidationError(
                        "نام و نام خانوادگی با سفارش مطابقت ندارد."
                    )

                order = (
                    Order.objects.select_for_update()
                    .select_related("user")
                    .get(pk=order.pk)
                )

                _apply_tracking_and_ship(
                    order=order,
                    tracking_code=tracking_code,
                    changed_by=changed_by,
                )
        except ValidationError as exc:
            detail = exc.detail
            if isinstance(detail, list):
                message = "; ".join(str(item) for item in detail)
            elif isinstance(detail, dict):
                message = "; ".join(
                    f"{key}: {value}" for key, value in detail.items()
                )
            else:
                message = str(detail)
            result.errors.append(
                RowError(
                    row_number=excel_row_index,
                    order_number=order_number,
                    message=message,
                )
            )
        except Exception as exc:
            result.errors.append(
                RowError(
                    row_number=excel_row_index,
                    order_number=order_number,
                    message=str(exc),
                )
            )
        else:
            result.success_count += 1

    return result


def build_empty_tracking_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "کد رهگیری"
    sheet.append(list(HEADERS))

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.append(
        [
            "ORD-1001",
            "علی",
            "محمدی",
            "1234567890",
            "09121234567",
        ]
    )

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 16

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_preparing_orders_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "آماده ارسال"
    sheet.append(list(HEADERS))

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    orders = (
        Order.objects.select_related("user")
        .filter(status=OrderStatus.PREPARING)
        .order_by("id")
    )

    for order in orders:
        user = order.user
        sheet.append(
            [
                _order_lookup_key(order),
                user.first_name or "",
                user.last_name or "",
                "",
                order.phone_number or user.phone_number or "",
            ]
        )

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 16

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
